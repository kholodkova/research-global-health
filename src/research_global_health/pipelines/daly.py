"""Extract a small local series from WHO GHE summary workbooks."""

import argparse
import csv
import math
from pathlib import Path

from openpyxl import load_workbook

YEARS = (2000, 2010, 2015, 2019, 2020, 2021)


def read_total(path: Path, year: int) -> float:
    """Read Persons / All Causes / RUS, in thousands of DALYs."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["All ages"]
        # Some WHO files declare 16,384 columns despite having far fewer cells.
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        header = [next(rows) for _ in range(8)]
        if "DALY ('000)" not in str(header[4][3]):
            raise ValueError(f"{path.name}: expected DALYs in thousands")
        if str(year) not in str(header[5][3]):
            raise ValueError(f"{path.name}: year does not match {year}")
        columns = [i for i, value in enumerate(header[7]) if value == "RUS"]
        if len(columns) != 1:
            raise ValueError(f"{path.name}: expected exactly one RUS column")
        column = columns[0]
        totals: list[float] = []
        for row in rows:
            if len(row) < 4 or row[0] != "Persons" or row[1] != 0:
                continue
            if row[3] != "All Causes":
                raise ValueError(f"{path.name}: unexpected cause label")
            value = row[column] if len(row) > column else None
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                raise ValueError(f"{path.name}: DALY total is not numeric")
            if not math.isfinite(value) or value <= 0:
                raise ValueError(f"{path.name}: DALY total must be finite and positive")
            totals.append(float(value))
        if len(totals) != 1:
            raise ValueError(
                f"{path.name}: expected exactly one Persons / All Causes row"
            )
        return totals[0]
    finally:
        workbook.close()


def extract(source_dir: Path, output: Path) -> None:
    """Validate all inputs before writing the local CSV."""
    totals = [
        read_total(source_dir / f"ghe2021_daly_bycountry_{year}.xlsx", year)
        for year in YEARS
    ]
    baseline = totals[0]
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            [
                "year",
                "country_iso3",
                "sex",
                "age",
                "ghe_code",
                "daly_thousands",
                "change_pct_vs_2000",
            ]
        )
        for year, total in zip(YEARS, totals, strict=True):
            writer.writerow(
                [
                    year,
                    "RUS",
                    "Persons",
                    "All ages",
                    0,
                    total,
                    (total / baseline - 1) * 100,
                ]
            )


def main() -> None:
    """Run the local extraction without publishing data."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=Path("data/local/daly_rus.csv"))
    args = parser.parse_args()
    extract(args.source_dir, args.output)
    print(f"Saved {len(YEARS)} rows to {args.output}")


if __name__ == "__main__":
    main()
