"""Small synthetic workbooks test extraction without redistributing WHO data."""

import csv
from pathlib import Path

import pytest
from openpyxl import Workbook

from research_global_health.pipelines.daly import YEARS, extract, read_total


def make_workbook(path: Path, year: int, value: float = 100.0) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "All ages"
    sheet.cell(5, 4, "Estimated DALY ('000) by cause, sex")
    sheet.cell(6, 4, f"and country or area (1), {year}")
    sheet.cell(8, 8, "USA")
    sheet.cell(8, 9, "RUS")
    for index, (sex, code, label, total) in enumerate(
        [
            ("Persons", 0, "All Causes", value),
            ("Male", 0, "All Causes", 50),
            ("Persons", 10, "Other cause", 25),
        ],
        start=11,
    ):
        sheet.cell(index, 1, sex)
        sheet.cell(index, 2, code)
        sheet.cell(index, 4, label)
        sheet.cell(index, 8, 999)
        sheet.cell(index, 9, total)
    workbook.save(path)
    workbook.close()


def test_selects_country_sex_and_total(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    make_workbook(path, 2021)
    assert read_total(path, 2021) == 100


@pytest.mark.parametrize("value", [0.0, -1.0])
def test_rejects_nonpositive_total(tmp_path: Path, value: float) -> None:
    path = tmp_path / "source.xlsx"
    make_workbook(path, 2021, value)
    with pytest.raises(ValueError, match="positive"):
        read_total(path, 2021)


def test_rejects_wrong_year(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    make_workbook(path, 2020)
    with pytest.raises(ValueError, match="year"):
        read_total(path, 2021)


def test_extract_changes_and_preserves_output_on_missing_input(tmp_path: Path) -> None:
    for i, year in enumerate(YEARS):
        make_workbook(
            tmp_path / f"ghe2021_daly_bycountry_{year}.xlsx", year, 100 + i * 10
        )
    output = tmp_path / "result.csv"
    extract(tmp_path, output)
    with output.open(newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 6
    assert float(rows[0]["change_pct_vs_2000"]) == 0
    assert float(rows[-1]["change_pct_vs_2000"]) == pytest.approx(50)
    original = output.read_bytes()
    (tmp_path / "ghe2021_daly_bycountry_2021.xlsx").unlink()
    with pytest.raises(FileNotFoundError):
        extract(tmp_path, output)
    assert output.read_bytes() == original
